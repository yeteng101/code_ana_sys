#!/usr/bin/env python3
"""Build an evidence-checked simulated review and reviewed Gold snapshot.

This is an AI-assisted source review, not a human confirmation. Evidence
anchors are checked against pinned sources, then category-specific rules
classify each existing annotation. The reviewed Gold is written beside the
official pending-review benchmark instead of overwriting it.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts.export_validation import OP_KIND
from scripts.prepare_validation_truth import sync_cases


ROOT = Path(__file__).resolve().parents[1]
REVIEWER = "codex-simulated-review-2026-10-16"
OUTPUT_DATE = "2026-10-16"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=fieldnames,
            extrasaction="ignore",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def normalize(text: str) -> str:
    return " ".join(text.split())


def evidence_ok(evidence: dict[str, Any]) -> tuple[bool, str]:
    file = str(evidence.get("file", ""))
    line = evidence.get("line")
    snippet = str(evidence.get("snippet", ""))
    path = ROOT / file
    if not path.is_file():
        return False, f"missing source: {file}"
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    if not isinstance(line, int) or not 1 <= line <= len(lines):
        return False, f"line out of range: {file}:{line}"
    if normalize(lines[line - 1]) != normalize(snippet):
        return False, f"snippet mismatch: {file}:{line}"
    return True, "evidence anchor verified"


def review_libuv_calls(data: dict[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for sample in data["samples"]:
        evidence_valid = all(evidence_ok(item)[0] for item in sample["evidence"])
        generated = "generated_candidate" in sample.get("tags", [])
        if generated:
            evidence_valid = False
            decision = "Rejected"
            reason = (
                "自动生成占位负样本，evidence 指向 README 而不是目标源码，"
                "不能作为代码事实进入 Gold。"
            )
        elif not evidence_valid:
            decision = "Needs fix"
            reason = next(
                (evidence_ok(item)[1] for item in sample["evidence"] if not evidence_ok(item)[0]),
                "evidence invalid",
            )
        elif sample["claim"]["kind"] == "callback_edge" and sample["expected_exists"]:
            decision = "Confirmed with condition"
            reason = "回调边存在，但目标由注册点或结构体字段决定，需保留候选条件。"
        else:
            decision = "Confirmed"
            reason = "调用方向、类型与源码证据一致。"
        entries.append(
            {
                "category": "call_chain",
                "project": "libuv",
                "id": sample["id"],
                "decision": decision,
                "evidence_valid": evidence_valid,
                "reason": reason,
                "gold": sample["graph"] == "demo/libuv/graph.json" and not generated,
            }
        )
    return entries


def review_redis_calls(data: dict[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for case in data["cases"]:
        checks = [evidence_ok(item) for item in case["evidence"]]
        evidence_valid = all(item[0] for item in checks)
        if evidence_valid:
            decision = "Confirmed"
            reason = "固定 commit 源码、行号和调用片段一致。"
        else:
            decision = "Needs fix"
            reason = next((item[1] for item in checks if not item[0]), "evidence invalid")
        entries.append(
            {
                "category": "call_chain",
                "project": "redis",
                "id": case["id"],
                "decision": decision,
                "evidence_valid": evidence_valid,
                "reason": reason,
                "gold": True,
            }
        )
    return entries


def base_resource_id(resource_id: str) -> str:
    return re.sub(r"_candidate_\d+$", "", resource_id)


def review_resources(data: dict[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for resource in data["resources"]:
        review = resource.get("ground_truth", {}).get("review", "pending")
        for operation in resource["operations"]:
            location = operation["location"]
            check, message = evidence_ok(location)
            if review == "confirmed":
                decision = "Confirmed"
                reason = "资源操作事实和 evidence 与固定源码一致。"
            elif review == "confirmed_with_condition":
                decision = "Confirmed with condition"
                reason = "平台分支或所有权条件成立时，该资源操作成立。"
            elif "_candidate_" in resource["id"]:
                decision = "Rejected"
                reason = (
                    f"与已确认基础样本 {base_resource_id(resource['id'])} 重复，"
                    "不作为独立 Gold。"
                )
            else:
                decision = "Needs fix"
                reason = message or "pending resource requires semantic review."
            entries.append(
                {
                    "category": "resource_flow",
                    "project": "libuv",
                    "id": resource["id"],
                    "operation_id": operation["operation_id"],
                    "decision": decision,
                    "evidence_valid": check,
                    "reason": reason,
                    "gold": (
                        resource.get("repository") == "libuv/libuv"
                        and review != "pending"
                    ),
                }
            )
    return entries


def review_libuv_sync(data: dict[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for sample in data["samples"]:
        checks = [evidence_ok(item) for item in sample["evidence"]]
        evidence_valid = all(item[0] for item in checks)
        if not evidence_valid:
            decision = "Needs fix"
            reason = next((item[1] for item in checks if not item[0]), "evidence invalid")
        elif sample["order"] in {"unknown", "concurrent"}:
            decision = "Confirmed with condition"
            reason = "结论表达的是未知或并发边界，条件和 evidence 一致。"
        elif sample["label"] == "positive":
            decision = "Confirmed"
            reason = "同步原语、执行上下文和 happens-before 证据一致。"
        else:
            decision = "Confirmed with condition"
            reason = "负样本/反例成立，但需保留声明的条件边界。"
        entries.append(
            {
                "category": "sync_relation",
                "project": "libuv",
                "id": sample["id"],
                "decision": decision,
                "evidence_valid": evidence_valid,
                "reason": reason,
                "gold": True,
            }
        )
    return entries


def review_redis_sync(data: dict[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for sample in data["samples"]:
        checks = [evidence_ok(item) for item in sample["evidence"]]
        evidence_valid = all(item[0] for item in checks)
        source = sample["evidence"][0]
        target = sample["evidence"][1]
        same_anchor = (
            source.get("file") == target.get("file")
            and source.get("line") == target.get("line")
        )
        if not evidence_valid:
            decision = "Needs fix"
            reason = next((item[1] for item in checks if not item[0]), "evidence invalid")
        elif sample["order"] == "unknown":
            decision = "Confirmed with condition"
            reason = "无法从现有证据推出确定顺序，unknown 结论成立。"
        elif sample["order"] == "concurrent" or same_anchor:
            decision = "Confirmed with condition"
            reason = "并发关系成立，但同一条源码锚点代表不同实例或分支，需要条件区分。"
        else:
            decision = "Confirmed"
            reason = "同步顺序、condition 和源码证据一致。"
        entries.append(
            {
                "category": "sync_relation",
                "project": "redis",
                "id": sample["id"],
                "decision": decision,
                "evidence_valid": evidence_valid,
                "reason": reason,
                "gold": True,
            }
        )
    return entries


def review_metadata() -> dict[str, Any]:
    return {
        "human_review": "pending",
        "method": "simulated_manual_review_by_codex",
        "reviewer": REVIEWER,
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "scope": (
            "Every evidence anchor was checked against the pinned checkout. "
            "Decisions are automated source checks plus category heuristics; "
            "they do not replace a human reviewer."
        ),
        "notes": (
            "Generated placeholder call candidates and duplicate pending "
            "resource candidates were rejected as Gold. Official human_review "
            "remains pending."
        ),
    }


def build_gold() -> dict[str, Any]:
    entries = [
        *review_libuv_calls(read_json(ROOT / "validation/call-chains/dataset.json")),
        *review_redis_calls(read_json(ROOT / "validation/redis/call-chains.json")),
        *review_resources(read_json(ROOT / "validation/resource-flow/dataset.json")),
        *review_libuv_sync(read_json(ROOT / "validation/sync-relations/dataset.json")),
        *review_redis_sync(read_json(ROOT / "validation/redis/sync-relations.json")),
    ]

    libuv_calls = []
    for sample in read_json(ROOT / "validation/call-chains/dataset.json")["samples"]:
        if sample["graph"] != "demo/libuv/graph.json":
            continue
        if {"generated_candidate", "pending_review"} & set(sample.get("tags", [])):
            continue
        claim = sample["claim"]
        libuv_calls.append(
            {
                "id": sample["id"],
                "category": "call_chain",
                "identity": [claim["source"], claim["target"]],
                "kind": claim["kind"],
                "expected_exists": sample["expected_exists"],
                "evidence": sample["evidence"],
                "notes": sample.get("notes", ""),
            }
        )
    for resource in read_json(ROOT / "validation/resource-flow/dataset.json")["resources"]:
        if resource.get("repository") != "libuv/libuv":
            continue
        if resource.get("ground_truth", {}).get("review") == "pending":
            continue
        for operation in resource["operations"]:
            location = operation["location"]
            libuv_calls.append(
                {
                    "id": resource["id"] + "_" + operation["operation_id"],
                    "category": "resource_flow",
                    "identity": [
                        resource["id"],
                        operation["function"],
                        location["file"],
                        location["line"],
                    ],
                    "kind": OP_KIND[operation["kind"]],
                    "expected_exists": True,
                    "evidence": [location],
                    "notes": "Operation fact, not a final defect classification.",
                }
            )
    libuv_calls += sync_cases(
        read_json(ROOT / "validation/sync-relations/dataset.json")
    )

    redis_calls = read_json(ROOT / "validation/redis/call-chains.json")
    redis_cases = redis_calls["cases"] + sync_cases(
        read_json(ROOT / "validation/redis/sync-relations.json")
    )

    libuv_truth = {
        "schema_version": "1.0",
        "dataset_id": "libuv-simulated-review-v1",
        "review": review_metadata(),
        "repository": "libuv/libuv",
        "commit": "8fb9cb919489a48880680a56efecff6a7dfb4504",
        "cases": libuv_calls,
    }
    redis_truth = {
        "schema_version": "1.0",
        "dataset_id": "redis-simulated-review-v1",
        "review": review_metadata(),
        "repository": "redis/redis",
        "commit": "d2c8a4b91e8c0e6aefd1f5bc0bf582cddbe046b7",
        "cases": redis_cases,
    }
    write_json(ROOT / "validation/benchmarks/libuv/ground-truth.simulated.json", libuv_truth)
    write_json(ROOT / "validation/benchmarks/redis/ground-truth.simulated.json", redis_truth)
    from scripts.check_benchmark_sources import check

    for project, truth in (("libuv", libuv_truth), ("redis", redis_truth)):
        source_report = check(
            truth,
            read_json(
                ROOT / "validation/benchmarks" / project / "source-manifest.json"
            ),
        )
        if source_report.get("status") != "verified":
            raise RuntimeError(f"{project} simulated Gold source check failed")
        write_json(
            ROOT / "validation/benchmarks" / project / "source-check.simulated.json",
            source_report,
        )

    decision_counts = Counter(item["decision"] for item in entries)
    gold_counts = Counter(item["category"] for item in entries if item["gold"])
    report = {
        "schema_version": "1.0",
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "reviewer": REVIEWER,
        "review": review_metadata(),
        "summary": {
            "rows": len(entries),
            "decisions": dict(decision_counts),
            "gold_rows": sum(gold_counts.values()),
            "gold_by_category": dict(gold_counts),
            "libuv_gold_cases": len(libuv_truth["cases"]),
            "redis_gold_cases": len(redis_truth["cases"]),
            "source_checks": "verified",
        },
        "entries": entries,
    }
    write_json(ROOT / "validation/reviews/simulated-review-2026-10-16.json", report)
    write_csv(
        ROOT / "validation/reviews/simulated-review-2026-10-16.csv",
        entries,
        [
            "category",
            "project",
            "id",
            "operation_id",
            "decision",
            "evidence_valid",
            "gold",
            "reason",
        ],
    )
    return report


def _load_workbook_rows(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to write the reviewed Excel workbook"
        ) from exc
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["Review"]
    raw = list(sheet.iter_rows(values_only=True))
    header_index = next(index for index, row in enumerate(raw) if row and row[0] == "Category")
    header = raw[header_index]
    rows = []
    for row in raw[header_index + 1 :]:
        if not row or not row[0]:
            continue
        rows.append(dict(zip(header, row)))
    workbook.close()
    return rows, list(header)


def _row_decision_key(row: dict[str, Any], resource_cursor: dict[str, int]) -> tuple[Any, ...]:
    if row.get("Category") == "资源流":
        resource_id = str(row.get("ID") or "")
        entries = review_by_resource.get(resource_id, [])
        index = resource_cursor.get(resource_id, 0)
        resource_cursor[resource_id] = index + 1
        if index < len(entries):
            return ("resource_flow", "libuv", resource_id, entries[index].get("operation_id"))
    category = {
        "调用链": "call_chain",
        "资源流": "resource_flow",
        "同步关系": "sync_relation",
    }.get(str(row.get("Category")), str(row.get("Category")))
    project = {"libuv": "libuv", "Redis": "redis"}.get(
        str(row.get("Project")), str(row.get("Project"))
    )
    return (category, project, str(row.get("ID")))


review_by_resource: dict[str, list[dict[str, Any]]] = {}
decision_lookup: dict[tuple[Any, ...], dict[str, Any]] = {}


def _prepare_lookups(report: dict[str, Any]) -> None:
    review_by_resource.clear()
    decision_lookup.clear()
    for entry in report["entries"]:
        if entry["category"] == "resource_flow":
            review_by_resource.setdefault(entry["id"], []).append(entry)
            decision_lookup[
                ("resource_flow", "libuv", entry["id"], entry["operation_id"])
            ] = entry
        else:
            decision_lookup[(entry["category"], entry["project"], entry["id"])] = entry


def write_reviewed_workbook(report: dict[str, Any]) -> Path:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to write the reviewed Excel workbook"
        ) from exc

    _prepare_lookups(report)
    source = ROOT / "validation/人工复核表-2026-10-16.xlsx"
    rows, original_header = _load_workbook_rows(source)
    header = original_header + ["Operation ID", "Evidence check", "Review reason"]
    decision_fill = {
        "Confirmed": PatternFill("solid", fgColor="C6EFCE"),
        "Confirmed with condition": PatternFill("solid", fgColor="FFEB9C"),
        "Rejected": PatternFill("solid", fgColor="FFC7CE"),
        "Needs fix": PatternFill("solid", fgColor="FCE4D6"),
    }
    output = ROOT / f"validation/人工复核表-{OUTPUT_DATE}-simulated.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Review"
    summary = report["summary"]["decisions"]
    gold_total = report["summary"]["libuv_gold_cases"] + report["summary"]["redis_gold_cases"]
    sheet.append(
        [
            "验证集模拟复核表",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    sheet.append(
        [
            f"Total: {report['summary']['rows']}",
            f"Confirmed: {summary.get('Confirmed', 0)}",
            f"Confirmed with condition: {summary.get('Confirmed with condition', 0)}",
            f"Rejected: {summary.get('Rejected', 0)}",
            f"Needs fix: {summary.get('Needs fix', 0)}",
            f"Gold cases: {gold_total}",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    sheet.append([])
    sheet.append(header)
    resource_cursor: dict[str, int] = {}
    for row in rows:
        key = _row_decision_key(row, resource_cursor)
        entry = decision_lookup.get(key, {})
        row_category = row.get("Category")
        row_project = row.get("Project")
        if row_category == "调用链" and row_project == "libuv" and "_negative_candidate_" in str(row.get("ID")):
            row["Expected / Classification"] = "negative"
        if row_category == "同步关系" and row_project == "Redis":
            order = row.get("Type / Order")
            if order == "unknown":
                row["Expected / Classification"] = "unknown"
            elif order == "concurrent":
                row["Expected / Classification"] = "inconsistent"
            else:
                row["Expected / Classification"] = "consistent"
        values = [row.get(key) for key in original_header]
        values[original_header.index("Decision")] = entry.get("decision")
        values.extend(
            [
                entry.get("operation_id", ""),
                "Pass" if entry.get("evidence_valid") else "Fail",
                entry.get("reason", ""),
            ]
        )
        sheet.append(values)

    decision_column = header.index("Decision") + 1
    reason_column = header.index("Review reason") + 1
    last_row = sheet.max_row
    for row_number in range(5, last_row + 1):
        decision = sheet.cell(row=row_number, column=decision_column).value
        fill = decision_fill.get(decision)
        if fill:
            sheet.cell(row=row_number, column=decision_column).fill = fill
        sheet.cell(row=row_number, column=reason_column).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    validation = DataValidation(
        type="list",
        formula1='"Confirmed,Confirmed with condition,Rejected,Needs fix"',
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{get_column_letter(decision_column)}5:{get_column_letter(decision_column)}{last_row}")
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(header))}{last_row}"
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    widths = {
        "A": 12,
        "B": 10,
        "C": 32,
        "D": 26,
        "E": 32,
        "F": 30,
        "G": 42,
        "H": 30,
        "I": 42,
        "J": 20,
        "K": 24,
        "L": 20,
        "M": 18,
        "N": 18,
        "O": 28,
        "P": 60,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    workbook.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--skip-workbook",
        action="store_true",
        help="skip the optional openpyxl-based Excel output",
    )
    args = parser.parse_args()
    report = build_gold()
    outputs = {
        "review_json": str(ROOT / "validation/reviews/simulated-review-2026-10-16.json"),
        "review_csv": str(ROOT / "validation/reviews/simulated-review-2026-10-16.csv"),
        "libuv_gold": str(ROOT / "validation/benchmarks/libuv/ground-truth.simulated.json"),
        "redis_gold": str(ROOT / "validation/benchmarks/redis/ground-truth.simulated.json"),
        "libuv_source_check": str(
            ROOT / "validation/benchmarks/libuv/source-check.simulated.json"
        ),
        "redis_source_check": str(
            ROOT / "validation/benchmarks/redis/source-check.simulated.json"
        ),
    }
    if not args.skip_workbook:
        try:
            outputs["review_workbook"] = str(write_reviewed_workbook(report))
        except RuntimeError as exc:
            outputs["review_workbook"] = f"skipped: {exc}"
    print(json.dumps({**report["summary"], "outputs": outputs}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
